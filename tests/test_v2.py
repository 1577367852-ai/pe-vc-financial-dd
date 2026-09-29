import json
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from common import ROOT,QA,dec,file_hash,write_text
from project import prepare,analyze,read_current,refresh,csv_load,csv_save,FILES,load_project
from fixtures_v2 import basic_project,source_book,confirm,replace_values
from ingest import extract
from cash import forecast_plan,build_events,opening_cash
from test_acceptance import observed


class Improvements(unittest.TestCase):
    def setUp(self):
        self.temp=Path(tempfile.mkdtemp(prefix=self._testMethodName+'-',dir=QA))

    def case(self,matrix=False):return basic_project(self.temp/'project',matrix=matrix)

    def test_T34_generic_identity(self):
        p=source_book(self.temp/'raw.xlsx',entity='合成另一个主体乙',period='2029-H1',cutoff='2029-06-30',matrix=True)
        f=self.temp/'case';prepare(f,[p])
        confirm(f,dict(entity='合成另一个主体乙',scope='合并',period='2029-H1',currency='CNY',cutoff='2029-06-30',baseline='2029-06-30',business_model='混合模式',horizon_months=2))
        r=analyze(f)
        self.assertEqual(r['name'],'合成另一个主体乙');self.assertEqual(r['cutoff'],'2029-06-30')
        self.assertEqual(r['metrics']['gross_margin']['value'],dec('.4'));self.assertEqual(r['cash']['rows'][-1]['date'],'2029-08-31')
        observed('T34',{'company':'乙','period':'2029-H1','layout':'多期间且科目B列'},r['metrics']['gross_margin'],'毛利率0.4，使用乙公司及2029期间')

    def test_T35_confirmation_and_local_missing(self):
        p=source_book(self.temp/'raw.xlsx');f=self.temp/'project';prepare(f,[p])
        with self.assertRaisesRegex(ValueError,'待确认'):analyze(f)
        self.assertEqual(json.loads((f/'results/manifest.json').read_text())['status'],'FAILED')
        confirm(f,dict(entity='合成新公司甲',scope='合并',period='2028-H1',currency='CNY',cutoff='2028-06-30',baseline='2028-06-30',business_model='产品销售型'))
        rows=csv_load(f/'科目确认.csv')
        for x in rows:
            if x['standard_account']=='net_profit':x['decision']='pending'
        csv_save(f/'科目确认.csv',FILES['科目确认.csv'],rows)
        r=analyze(f);self.assertEqual(r['metrics']['net_margin']['status'],'BLOCKED_INPUT');self.assertEqual(r['metrics']['gross_margin']['value'],dec('.4'))
        observed('T35','项目未确认；随后仅净利润待确认',{'net_margin':r['metrics']['net_margin']['status'],'gross_margin':'.4'},'未确认拒绝；单项缺失不阻断其他指标')

    def test_T36_integrated_bridges(self):
        f,p=self.case();r=analyze(f)
        expected={'QoE':110000,'NWC':230000,'NetDebt':-220000}
        for k,v in expected.items():self.assertEqual(r['bridges'][k]['value'],v)
        self.assertEqual(r['adjustments'][-1]['status'],'unverified')
        from openpyxl import load_workbook
        w=load_workbook(f/'results/财务分析.xlsx',data_only=False);self.assertIn('调整桥',w.sheetnames);self.assertIn('尽调问题',w.sheetnames);w.close()
        for k in expected:self.assertIn('bridge:'+k,read_current(f))
        observed('T36',{'net_profit':100000,'nwc':250000,'netdebt':-250000,'adjustments':[10000,-20000,30000],'unverified':90000},expected,expected)

    def test_T37_cross_bucket_collision(self):
        f,p=self.case();rows=csv_load(f/'调整事项.csv');rows[1]['economic_item_id']=rows[0]['economic_item_id']
        csv_save(f/'调整事项.csv',FILES['调整事项.csv'],rows);r=analyze(f)
        self.assertEqual(r['bridges']['QoE']['status'],'BLOCKED_INPUT');self.assertEqual(r['bridges']['NWC']['status'],'BLOCKED_INPUT');self.assertEqual(r['bridges']['NetDebt']['value'],-220000)
        observed('T37','QoE和NWC同经济事项编号',{k:v['status'] for k,v in r['bridges'].items()},'仅冲突桶阻断，净债务仍可算')

    def test_T38_evidence_and_duplicate_adjustments(self):
        f,p=self.case();rows=csv_load(f/'调整事项.csv');rows[0]['evidence_sha256']='expired'
        csv_save(f/'调整事项.csv',FILES['调整事项.csv'],rows);r=analyze(f);self.assertEqual(r['bridges']['QoE']['value'],100000)
        rows.append(dict(rows[0]));csv_save(f/'调整事项.csv',FILES['调整事项.csv'],rows)
        with self.assertRaisesRegex(ValueError,'重复调整'):analyze(f)
        observed('T38','过期证据及重复调整ID',{'stale_excluded':100000,'duplicate':'FAILED'},'无依据不采纳；重复ID拒绝')

    def test_T39_dynamic_findings(self):
        f,p=self.case();r=analyze(f);gap=next(x for x in r['findings'] if x['finding_id']=='cash-gap')
        self.assertEqual(gap['amount'],150000);self.assertTrue(gap['source']);self.assertIn('付款',gap['request'])
        csv_save(f/'持续费用.csv',FILES['持续费用.csv'],[dict(id='run',amount=100000,source='合成修改预算')])
        r=analyze(f);self.assertFalse(any(x['finding_id']=='cash-gap' for x in r['findings']))
        observed('T39','月费用20万改10万',{'gap_before':150000,'gap_after':0},'风险发现随计算结果消失，非固定模板')

    def test_T40_multiperiod_units_and_formula(self):
        from openpyxl import load_workbook
        p=source_book(self.temp/'matrix.xlsx',matrix=True)
        w=load_workbook(p);s=w.active
        for row in s:
            if row[1].value=='营业收入':row[2].value=80;row[3].value='=40+60';row[4].value='万元'
        w.save(p);w.close();r=extract(p)
        revenue=[x for x in r['records'] if x['standard_account']=='revenue']
        self.assertEqual(len(revenue),2);self.assertEqual([x['value'] for x in revenue],[800000,1000000]);self.assertEqual(revenue[1]['status'],'RECOMPUTED')
        self.assertEqual(revenue[1]['source']['cell'],'D10')
        observed('T40','B科目,C2027,D2028,E单位；100万元公式',[x['value'] for x in revenue],[800000,1000000])

    def test_T41_multipage_pdf(self):
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from examples import FONT
        pdfmetrics.registerFont(TTFont('V2CJK',FONT));p=self.temp/'crosspage.pdf';c=canvas.Canvas(str(p))
        lines=['合成样本','主体: 合成甲','期间: 2028年度','范围: 合并','币种: CNY','版本: v1','截止日: 2028-12-31','单位: 万元','科目 | 2027年度 | 2028年度 | 单位','营业收入 | 80 | 100 | 万元']
        c.setFont('V2CJK',12)
        for i,line in enumerate(lines):c.drawString(30,780-i*30,line)
        c.showPage();c.setFont('V2CJK',12);c.drawString(30,780,'营业成本 | 50 | 60 | 万元');c.save()
        r=extract(p);cost=[x for x in r['records'] if x['standard_account']=='cogs']
        self.assertEqual(len(cost),2);self.assertEqual(cost[1]['period'],'2028年度');self.assertIsNone(cost[1]['value'])
        self.assertEqual(cost[1]['candidate_value'],600000);self.assertEqual(cost[1]['source']['physical_page'],2);self.assertTrue(cost[1]['source']['bbox'])
        observed('T41','跨页两期报表、第二页不重复主体',cost,'保留两期及物理页；跨页继承需确认，60万元不能直接采用')

    def test_T42_noisy_ocr_review(self):
        from PIL import Image,ImageFilter
        original=ROOT/'examples/biotech/合成现金截图.png';p=self.temp/'noisy.png'
        with Image.open(original) as im:im.resize((960,492)).filter(ImageFilter.GaussianBlur(.6)).save(p)
        r=extract(p);self.assertEqual(r['status'],'NEEDS_REVIEW');self.assertTrue(r['raw']['text'])
        self.assertTrue(all(x['value'] is None for x in r['records']))
        observed('T42',{'image':'合成缩小模糊截图'},dict(candidate_count=len(r['records']),auto_adopted=0),'实际本地OCR运行，所有候选仍需复核')

    def test_T43_input_change_refresh(self):
        from openpyxl import load_workbook
        f,p=self.case();r=analyze(f);w=load_workbook(p)
        for row in w.active:
            if row[0].value=='货币资金':row[1].value=600000
        w.save(p);w.close()
        with self.assertRaisesRegex(ValueError,'STALE'):read_current(f)
        with self.assertRaisesRegex(ValueError,'过期'):analyze(f)
        refresh(f);self.assertTrue(all(x['decision']=='pending' for x in csv_load(f/'科目确认.csv')))
        confirm(f,dict(entity='合成新公司甲',scope='合并',period='2028-H1',currency='CNY',cutoff='2028-06-30',baseline='2028-06-30',business_model='产品销售型'))
        r=analyze(f);self.assertEqual(r['opening']['value'],550000);self.assertEqual(r['bridges']['QoE']['value'],100000)
        self.assertIn('550000',read_current(f))
        observed('T43','原件现金50万改60万',{'opening':550000,'old_adjustment_excluded':True},'读旧结果拒绝；refresh后重新确认；旧调整依据失效')

    def test_T44_assumption_update_sync(self):
        f,p=self.case();a=analyze(f);replace_values(f,'情景设置.csv',{'delay_months':3})
        with self.assertRaisesRegex(ValueError,'STALE'):read_current(f)
        b=analyze(f);self.assertEqual(len(b['cash']['rows']),6);self.assertEqual(b['cash']['max_zero_gap'],750000)
        saved=json.loads((f/'results/result.json').read_text());self.assertEqual(dec(saved['cash']['max_zero_gap']),750000)
        self.assertIn('750000',read_current(f));self.assertIn('750000',read_current(f,'learning'))
        observed('T44','延迟0→3月',{'rows':6,'gap':750000},'表格和两份摘要统一来自新计算')

    def test_T45_contract_version(self):
        f,p=self.case();contract=dict(contract_id='C1',version='v1',active='yes',total=10000,paid=0,obligation_id='C1-A',amount=10000,date='2028-07-20',timing='fixed',source='合成合同')
        csv_save(f/'合同节点.csv',FILES['合同节点.csv'],[contract]);analyze(f)
        contract['version']='v2';contract['total']=20000;contract['amount']=20000
        csv_save(f/'合同节点.csv',FILES['合同节点.csv'],[contract])
        with self.assertRaisesRegex(ValueError,'STALE'):read_current(f)
        r=analyze(f);e=next(x for x in r['cash']['schedule']['events'] if x['id']=='C1-A')
        self.assertEqual(e['contract_version'],'v2');self.assertEqual(e['amount'],20000)
        observed('T45','同合同新版1万→2万',e,'旧结果拒绝；新版金额2万仅计一次')

    def test_T46_code_change_and_output_tamper(self):
        f,p=self.case();analyze(f)
        with patch('reporting.code_hash',return_value='changed'):
            with self.assertRaisesRegex(ValueError,'STALE'):read_current(f)
        analyze(f);write_text(f/'results/工作摘要.md','伪造旧摘要')
        with self.assertRaisesRegex(ValueError,'STALE'):read_current(f)
        analyze(f);self.assertNotIn('伪造',read_current(f))
        observed('T46','模拟代码指纹变化；篡改输出',{'rejected':2,'recomputed':True},'两者均拒绝，重新计算恢复')

    def test_T47_failed_run_invalidates_old(self):
        f,p=self.case();analyze(f);replace_values(f,'项目设置.csv',{'confirmed_by':''})
        with self.assertRaises(ValueError):analyze(f)
        self.assertEqual(json.loads((f/'results/manifest.json').read_text())['status'],'FAILED')
        with self.assertRaises(ValueError):read_current(f)
        observed('T47','新输入失败，目录中仍有旧表格',{'manifest':'FAILED','view':'REFUSED'},'旧结果不再READY且不能通过view读取')

    def plan(self,baseline,amount,months=2,until=None):
        recurring=dict(id='r',amount=amount,source='合成预算',proration_source='合成合同：按自然日计费，月末支付')
        if until:recurring['until_milestone']='m'
        return dict(baseline=baseline,horizon_months=months,opening=opening_cash(1000000,0,baseline,baseline),
                    partial_month_policy='actual_days_month_end',recurring=[recurring],milestones={'m':until} if until else {})

    def test_T48_partial_start(self):
        r=forecast_plan(self.plan('2028-01-15',31000),{})
        self.assertEqual([x['outflows'] for x in r['rows']],[16000,31000]);self.assertEqual(r['rows'][0]['date'],'2028-01-31')
        observed('T48','1月15日基准，3.1万元月费，2个期间',[x['outflows'] for x in r['rows']],[16000,31000])

    def test_T49_leap_and_end(self):
        a=forecast_plan(self.plan('2028-02-14',2900,1),{});b=forecast_plan(self.plan('2027-02-14',2800,1),{})
        c=forecast_plan(self.plan('2028-01-31',2900,2,'2028-02-15'),{})
        self.assertEqual(a['rows'][0]['outflows'],1500);self.assertEqual(b['rows'][0]['outflows'],1400);self.assertEqual(c['rows'][0]['outflows'],1500)
        self.assertEqual(c['rows'][1]['outflows'],0)
        observed('T49','闰年/非闰年/月中停止',[a['rows'][0]['outflows'],b['rows'][0]['outflows'],c['rows'][0]['outflows']],[1500,1400,1500])

    def test_T50_partial_delay_linkage(self):
        p=self.plan('2028-01-15',3100,3,'2028-03-15')
        p['events']=[dict(id='fixed',date='2028-02-10',amount=1000,direction='out',category='opex',source='合成固定合同'),
                     dict(id='node',timing='milestone',milestone='m',amount=5000,direction='out',category='opex',source='合成节点合同'),
                     dict(id='F',timing='milestone',milestone='m',amount=20000,direction='in',category='primary_financing',condition='data',source='合成融资意向')]
        a=forecast_plan(p,dict(include_financing=True,conditions_met=['data']));b=forecast_plan(p,dict(delay_months=3,include_financing=True,conditions_met=['data']))
        events=b['schedule']['events'];byid={e['id']:e for e in events}
        self.assertEqual(byid['fixed']['date'],'2028-02-10');self.assertEqual(byid['node']['date'],'2028-06-15');self.assertEqual(byid['F']['date'],'2028-06-15')
        self.assertEqual(sum(e['amount'] for e in a['schedule']['events'] if e['origin']=='recurring'),6200)
        self.assertEqual(sum(e['amount'] for e in events if e['origin']=='recurring'),15550)
        self.assertEqual(len(b['rows']),6)
        observed('T50','月中基准，3月15日节点延至6月15日',{'base_recurring':6200,'delayed_recurring':15550,'fixed':byid['fixed']['date'],'node':byid['node']['date']},'持续费补全；节点和融资后移；固定日期不动')

    def test_T51_partial_requires_basis(self):
        p=self.plan('2028-01-15',3100);del p['recurring'][0]['proration_source']
        self.assertEqual(forecast_plan(p,{})['status'],'BLOCKED_INPUT')
        observed('T51','月中费用无分摊依据','BLOCKED_INPUT','不能默认按天拆分')

    def test_T52_manual_unknown_account(self):
        from openpyxl import load_workbook
        f,p=self.case();w=load_workbook(p);w.active['G1']='未识别净收益';w.active['H1']=12345;w.save(p);w.close()
        refresh(f);confirm(f,dict(entity='合成新公司甲',scope='合并',period='2028-H1',currency='CNY',cutoff='2028-06-30',baseline='2028-06-30',business_model='产品销售型'))
        rows=csv_load(f/'科目确认.csv')
        for x in rows:
            if x['standard_account']=='net_profit':x['decision']='exclude'
        csv_save(f/'科目确认.csv',FILES['科目确认.csv'],rows)
        m=dict(file=str(p),sha256=file_hash(p),sheet='不同版式报表',cell='H1',original_account='未识别净收益',standard_account='net_profit',
               entity='合成新公司甲',scope='合并',period='2028-H1',currency='CNY',version='v2',cutoff='2028-06-30',unit='元',confirmed_raw=12345,reviewer='合成确认人',evidence='人工定位H1')
        csv_save(f/'补充映射.csv',FILES['补充映射.csv'],[m]);r=analyze(f);self.assertEqual(r['metrics']['net_margin']['value'],dec('.012345'))
        observed('T52','未知科目H1，明确人工映射',r['metrics']['net_margin']['value'],'.012345；位置与原值保留')

    def test_T53_reviewed_screenshot(self):
        f,p=self.case();img=ROOT/'examples/biotech/合成现金截图.png'
        csv_save(f/'材料清单.csv',FILES['材料清单.csv'],[{'path':str(p)},{'path':str(img)}])
        # Attach an image as supplementary evidence, not silently replace another entity's cash.
        m=dict(file=str(img),sha256=file_hash(img),bbox='0,0,1600,820',original_account='货币资金',standard_account='book_cash',
               entity='合成研发药企',scope='合并',period='2026-H1',currency='CNY',version='synthetic-v1',cutoff='2026-06-30',unit='元',confirmed_raw=120000000,reviewer='合成确认人',evidence='按生成规格复核图中货币资金')
        csv_save(f/'补充映射.csv',FILES['补充映射.csv'],[m]);r=analyze(f)
        manual=[x for x in r['records'] if x.get('review',{}).get('bbox')]
        self.assertEqual(manual[0]['value'],120000000);self.assertEqual(r['opening']['value'],450000)
        observed('T53','补充另一主体截图确认',{'image_value':120000000,'selected_company_cash':450000},'确认保留；不同主体不得混入')

    def test_T54_original_v1_unchanged(self):
        manifest=json.loads((ROOT/'examples/原始样本指纹.json').read_text())
        for rel,h in manifest.items():self.assertEqual(file_hash(ROOT/rel),h)
        observed('T54','公开仓库原始合成材料指纹',{'verified_files':len(manifest)},'原始材料指纹与清单一致')

    def test_T55_csv_injection(self):
        f=self.temp/'safe.csv';csv_save(f,['text','amount'],[dict(text='=WEBSERVICE("https://example.invalid")',amount=-123)])
        text=f.read_text(encoding='utf-8-sig')
        self.assertIn("'=WEBSERVICE",text);self.assertEqual(csv_load(f)[0]['amount'],'-123')
        observed('T55','CSV中外链公式文字和负数',{'formula_escaped':True,'negative_preserved':True},'公式转文字，金额负数保留')

    def test_T56_midrun_mutation(self):
        f,p=self.case()
        original=load_project
        def mutate(folder):
            loaded=original(folder)
            replace_values(f,'情景设置.csv',{'delay_months':3})
            return loaded
        with patch('project.load_project',side_effect=mutate):
            with self.assertRaisesRegex(ValueError,'运行期间'):analyze(f)
        self.assertEqual(json.loads((f/'results/manifest.json').read_text())['status'],'FAILED')
        observed('T56','计算中途改变假设',{'publication':'REFUSED'},'拒绝发布不同版本混合结果')

    def test_T57_saved_types(self):
        from openpyxl import load_workbook
        from datetime import datetime
        f,p=self.case();analyze(f);w=load_workbook(f/'results/财务分析.xlsx')
        self.assertEqual(w['调整与问题']['D5'].value,-20000);self.assertIsInstance(w['调整与问题']['D5'].value,(int,float))
        self.assertIsInstance(w['现金预测']['A4'].value,datetime)
        self.assertEqual(w['指标与校验']['C5'].number_format,'0.0%')
        w.close();observed('T57','保存后的Excel负调整/日期/比率',{'negative':-20000,'date_type':'datetime','ratio_format':'0.0%'},'金额日期保持数值型，比例按百分比显示')


if __name__=='__main__':unittest.main()
