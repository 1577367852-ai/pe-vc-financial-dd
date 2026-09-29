"""Clearly synthetic fixtures and review actions for regression, never real evidence."""
from pathlib import Path
from copy import deepcopy
from common import ROOT, QA, allowed_write, file_hash, write_json
from project import prepare, csv_save, csv_load, FILES


def replace_values(folder, name, updates):
    rows=csv_load(folder/name)
    bykey={r['key']:r for r in rows}
    for key,value in updates.items():bykey[key]=dict(key=key,value=value,source='合成确认依据；不代表真实项目')
    csv_save(folder/name,FILES[name],list(bykey.values()))


def source_book(path, entity='合成新公司甲', period='2028-H1', cutoff='2028-06-30', matrix=False, revenue=1000000):
    from openpyxl import Workbook
    p=allowed_write(path);p.parent.mkdir(parents=True,exist_ok=True)
    w=Workbook();s=w.active;s.title='不同版式报表'
    meta={'主体':entity,'期间':period,'截止日':cutoff,'范围':'合并','币种':'CNY','版本':'synthetic-v2','单位':'元'}
    s.append(['合成样本，无真实公司信息'])
    for k,v in meta.items():s.append([k,v])
    facts={'营业收入':revenue,'营业成本':600000,'净利润':100000,'资产总计':2000000,'负债合计':500000,
           '所有者权益合计':1500000,'货币资金':500000,'受限现金':50000,'有息债务':200000,
           '应收账款':200000,'存货':150000,'其他经营性流动资产':30000,'应付账款':120000,'其他经营性流动负债':10000,
           '经营现金流出':600000,'经营现金流入':400000,'期间月数':6,'经营活动现金流量净额':-200000,'现金资本支出':30000}
    if matrix:
        s.append([None,'项目','2027-H1',period,'单位'])
        for k,v in facts.items():s.append([None,k,v,v,'月' if k=='期间月数' else '元'])
    else:
        s.append(['科目','金额','单位'])
        for k,v in facts.items():s.append([k,v,'月' if k=='期间月数' else '元'])
    e=w.create_sheet('合成凭证');e.append(['编号','科目','金额','说明'])
    e.append(['Q1','一次性已税后费用',10000,'仅供算例，批准加回'])
    e.append(['N1','滞销存货减值',-20000,'仅供算例，调整NWC'])
    e.append(['D1','债务性项目',30000,'仅供算例，净债务加项'])
    w.save(p);w.close();return p


def confirm(folder, config):
    replace_values(folder,'项目设置.csv',dict(config,confirmed_by='合成测试确认人',confirmation_evidence='按样本生成规格核对；非人工审计',synthetic_only='yes'))
    rows=csv_load(folder/'科目确认.csv')
    for r in rows:
        if r.get('period')!=config['period']:r['decision']='exclude';continue
        r.update(decision='include',reviewer='合成测试确认人',evidence='样本预设真值',cutoff=config['cutoff'])
    csv_save(folder/'科目确认.csv',FILES['科目确认.csv'],rows)


def basic_project(folder, source=None, matrix=False):
    f=allowed_write(folder)
    p=source or source_book(f.parent/(f.name+'-原件.xlsx'),matrix=matrix)
    prepare(f,[p])
    cfg=dict(entity='合成新公司甲',period='2028-H1',scope='合并',currency='CNY',cutoff='2028-06-30',baseline='2028-06-30',
             business_model='产品销售型',horizon_months='3',minimum='100000')
    confirm(f,cfg)
    csv_save(f/'持续费用.csv',FILES['持续费用.csv'],[dict(id='run',amount=200000,source='合成月预算，不含合同节点')])
    csv_save(f/'调整基数.csv',FILES['调整基数.csv'],[
        dict(bucket='QoE',base_account='net_profit',source='合成：净利润税后口径'),
        dict(bucket='NWC',base_account='operating_nwc',source='合成：经营流动资产减经营流动负债'),
        dict(bucket='NetDebt',base_account='net_debt',source='合成：有息债务减同日非受限现金')])
    items=[]
    for i,(bucket,amount,reason) in enumerate([('QoE',10000,'一次性已税后费用加回'),('NWC',-20000,'滞销存货减值'),('NetDebt',30000,'债务性项目纳入')],2):
        items.append(dict(adjustment_id='A'+str(i),economic_item_id='EC'+str(i),bucket=bucket,amount=amount,status='supported',approved='yes',
                          reviewer='合成确认人',reason=reason,evidence_file=str(p),evidence_sha256=file_hash(p),evidence_location=f'合成凭证!C{i}'))
    items.append(dict(adjustment_id='UNVERIFIED',economic_item_id='EC9',bucket='QoE',amount=90000,status='unverified',approved='no',reason='未经核实的管理层主张'))
    csv_save(f/'调整事项.csv',FILES['调整事项.csv'],items)
    return f,p


def legacy_to_project(kind, folder):
    """Only a fixture converter knows legacy names; the actual analysis entry is generic."""
    from openpyxl import load_workbook
    from pipeline import table
    p=ROOT/'examples'/kind/'合成原始材料.xlsx'
    prepare(folder,[p]);w=load_workbook(p,data_only=True)
    params={r['参数']:r['值'] for r in table(w['分析参数'])}
    confirm(folder,dict(entity='合成研发药企' if kind=='biotech' else '合成医疗器械公司',period='2026-H1',scope='合并',currency='CNY',
        cutoff='2026-06-30',baseline=params['baseline'],business_model=params['business_model'],horizon_months=12,minimum=params['minimum'],
        complete_bridge='yes',coverage_source='合成样本完整性声明'))
    for src,dst in [('期间现金变动','现金变动.csv'),('持续预算','持续费用.csv'),('现金事件','现金事件.csv'),('合同节点','合同节点.csv')]:
        rows=table(w[src])
        if src=='合同节点':rows=[dict(r,active='yes') for r in rows]
        csv_save(folder/dst,FILES[dst],rows)
    csv_save(folder/'里程碑.csv',FILES['里程碑.csv'],[dict(id='readout',date=params['milestone'],source='合成预算里程碑')])
    csv_save(folder/'调整基数.csv',FILES['调整基数.csv'],[dict(bucket='QoE',base_account='net_profit',source='合成净利润口径'),dict(bucket='NWC',base_account='operating_nwc',source='合成经营范围')])
    w.close();return folder


def build_delivery():
    from project import analyze
    root=ROOT/'examples'/'v2';root.mkdir(parents=True,exist_ok=True)
    done=[]
    for kind in ('biotech','sales'):
        f=root/kind
        if not f.exists():legacy_to_project(kind,f)
        for name,fields in FILES.items():
            if not (f/name).exists():csv_save(f/name,fields,[])
        report=analyze(f);done.append(dict(project=str(f),cash_status=report['cash']['status']))
    f=root/'多期间及调整示例'
    if not f.exists():basic_project(f,matrix=True)
    for name,fields in FILES.items():
        if not (f/name).exists():csv_save(f/name,fields,[])
    # Rebuild confirmation stamps for this known synthetic fixture only.
    rows=csv_load(f/'调整事项.csv')
    for row in rows:
        if row.get('status')=='supported':row['evidence_sha256']=file_hash(Path(row['evidence_file']))
    csv_save(f/'调整事项.csv',FILES['调整事项.csv'],rows)
    report=analyze(f);done.append(dict(project=str(f),bridges=report['bridges']))
    write_json(ROOT/'outputs'/'v0.2示例清单.json',done)
    return done


if __name__=='__main__':build_delivery()
