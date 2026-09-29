"""Generate ONLY clearly labelled synthetic fixtures. No real-project data."""
from pathlib import Path
from common import ROOT, QA, allowed_write, file_hash, write_json

FONT='/System/Library/Fonts/Supplemental/Arial Unicode.ttf'


def metadata(name):
    return [('主体',name),('期间','2026-H1'),('范围','合并'),('币种','CNY'),('版本','synthetic-v1'),('截止日','2026-06-30'),('单位','元')]


def build_examples():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from PIL import Image,ImageDraw,ImageFont
    pdfmetrics.registerFont(TTFont('SyntheticCJK',FONT))
    targets=[]
    for kind in ('biotech','sales'):
        folder=allowed_write(ROOT/'examples'/kind);folder.mkdir(parents=True,exist_ok=True)
        name='合成研发药企' if kind=='biotech' else '合成医疗器械公司'
        if kind=='biotech':
            facts={'营业收入':0,'营业成本':0,'净利润':-50000000,'研发费用':40000000,'经营费用合计':50000000,
                   '资产总计':150000000,'负债合计':30000000,'所有者权益合计':120000000,
                   '流动资产合计':130000000,'流动负债合计':20000000,'货币资金':120000000,'受限现金':20000000,
                   '应收账款':0,'存货':0,'应付账款':5000000,'其他经营性流动资产':5000000,'其他经营性流动负债':0,
                   '经营活动现金流量净额':-48000000,'投资活动现金流量净额':-2000000,'筹资活动现金流量净额':30000000,
                   '期初现金及现金等价物':120000000,'期末现金及现金等价物':100000000,'汇率变动影响':0,
                   '现金资本支出':2000000,'经营现金流出':48000000,'经营现金流入':0,'期间月数':6}
        else:
            facts={'营业收入':100000000,'营业成本':60000000,'净利润':10000000,'研发费用':5000000,'经营费用合计':25000000,
                   '资产总计':100000000,'负债合计':40000000,'所有者权益合计':60000000,
                   '流动资产合计':60000000,'流动负债合计':25000000,'货币资金':20000000,'受限现金':2000000,
                   '应收账款':20000000,'存货':15000000,'应付账款':12000000,'其他经营性流动资产':3000000,'其他经营性流动负债':3000000,
                   '经营活动现金流量净额':12000000,'投资活动现金流量净额':-8000000,'筹资活动现金流量净额':-6000000,
                   '期初现金及现金等价物':20000000,'期末现金及现金等价物':18000000,'汇率变动影响':0,
                   '现金资本支出':8000000,'经营现金流出':88000000,'经营现金流入':100000000,'期间月数':6,
                   '期初权益':50000000,'期初应收账款':16000000,'期初存货':10000000,'期初应付账款':10000000,
                   '赊销收入':80000000,'赊购额':65000000,'期间天数':181}
        wb=Workbook();ws=wb.active;ws.title='财务表'
        ws.append(['合成样本，不对应真实公司'])
        for r in metadata(name):ws.append(list(r))
        ws.append(['科目','金额','单位'])
        for k,v in facts.items():ws.append([k,v,'月' if k=='期间月数' else ('天' if k=='期间天数' else '元')])
        ws.column_dimensions['A'].width=38;ws.column_dimensions['B'].width=24;ws.column_dimensions['C'].width=15
        ws.freeze_panes='B10';ws.sheet_view.showGridLines=False
        for c in ws[9]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='29475F')
        # Readable parameters and event schedules, not hand-authored user JSON.
        p=wb.create_sheet('分析参数');p.append(['参数','值','依据'])
        params=[('baseline','2026-09-30'),('complete_bridge',True),('minimum',30000000 if kind=='biotech' else 3000000),
                ('horizon_months',12),('milestone', '2027-09-30'),('business_model','研发型' if kind=='biotech' else '产品销售型')]
        for k,v in params:p.append([k,v,'合成管理预算v1；仅演示'])
        b=wb.create_sheet('期间现金变动');b.append(['id','date','amount','status','source'])
        b.append(['B001','2026-07-31',-10000000 if kind=='biotech' else -2000000,'actual','合成银行流水B001'])
        if kind=='biotech':b.append(['B002','2026-08-31',30000000,'actual','合成银行流水B002：已到账增资'])
        r=wb.create_sheet('持续预算');r.append(['id','amount','until_milestone','source'])
        r.append(['RUN',15000000 if kind=='biotech' else 4000000,'readout','合成预算第1行；不含合同节点款'])
        e=wb.create_sheet('现金事件');e.append(['id','amount','date','direction','category','timing','milestone','condition','source'])
        if kind=='biotech':
            e.append(['F001',60000000,None,'in','primary_financing','milestone','readout','data_available','合成意向融资；非承诺到账'])
            e.append(['OLD',10000000,'2026-12-31','in','secondary_sale','fixed',None,None,'合成老股交易'])
        else:
            from common import month_ends
            for i,d in enumerate(month_ends('2026-09-30',12),1):e.append([f'COL{i}',5000000,str(d),'in','operating_receipt','fixed',None,None,'合成回款预算；非确定收款'])
            e.append(['CAPEX',2000000,'2026-12-31','out','capex','fixed',None,None,'合成扩产预算'])
        c=wb.create_sheet('合同节点');c.append(['contract_id','version','total','paid','obligation_id','amount','date','timing','milestone','source'])
        if kind=='biotech':
            c.append(['CRO01','v1',30000000,5000000,'CRO01-A',5000000,'2026-12-31','fixed',None,'合成CRO合同第3条：固定付款'])
            c.append(['CRO01','v1',30000000,5000000,'CRO01-B',20000000,None,'milestone','readout','合成CRO合同第4条：数据读出付款'])
        for sh in list(wb)[1:]:
            sh.freeze_panes='B2';sh.sheet_view.showGridLines=False
            for col in range(1,sh.max_column+1):sh.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=24
        xlsx=folder/'合成原始材料.xlsx';wb.save(xlsx);wb.close();targets.append(xlsx)
        pdf=folder/'合成报表摘录.pdf';can=canvas.Canvas(str(pdf),pagesize=(595,842))
        lines=['合成样本，不对应真实公司']+[f'{k}: {v}' for k,v in metadata(name)]
        lines += [f'{k} | {v} | 元' for k,v in list(facts.items())[:10]]
        can.setFont('SyntheticCJK',12)
        for i,line in enumerate(lines):can.drawString(36,800-30*i,line)
        can.save();targets.append(pdf)
        img=Image.new('RGB',(1600,820),'white');draw=ImageDraw.Draw(img);font=ImageFont.truetype(FONT,38)
        img_lines=['合成样本，仅测试OCR']+[f'{k}: {v}' for k,v in metadata(name)] + [f'货币资金 | {facts["货币资金"]} | 元',f'受限现金 | {facts["受限现金"]} | 元']
        for i,line in enumerate(img_lines):draw.text((35,25+i*70),line,font=font,fill='black')
        png=folder/'合成现金截图.png';img.save(png);targets.append(png)
    write_json(ROOT/'examples'/'原始样本指纹.json',{str(p.relative_to(ROOT)):file_hash(p) for p in targets})


if __name__=='__main__':build_examples()
