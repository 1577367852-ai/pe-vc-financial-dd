"""Read-only extraction. Limited formulas are recomputed without eval or Excel execution."""
import ast
import csv
import io
import re
import subprocess
from pathlib import Path
from common import QA, allowed_write, dec, file_hash, fingerprint, result

ALIASES = {
    '营业收入': 'revenue', '营业成本': 'cogs', '净利润': 'net_profit', '研发费用': 'rd_expense',
    '经营费用合计': 'total_opex', '资产总计': 'assets', '负债合计': 'liabilities', '所有者权益合计': 'equity',
    '流动资产合计': 'current_assets', '流动负债合计': 'current_liabilities', '货币资金': 'book_cash',
    '受限现金': 'restricted_cash', '应收账款': 'ar', '存货': 'inventory', '应付账款': 'ap',
    '其他经营性流动资产': 'other_operating_ca', '其他经营性流动负债': 'other_operating_cl',
    '经营活动现金流量净额': 'ocf', '投资活动现金流量净额': 'icf', '筹资活动现金流量净额': 'fcf_financing',
    '期初现金及现金等价物': 'cash_open', '期末现金及现金等价物': 'cash_close', '汇率变动影响': 'fx_cash',
    '现金资本支出': 'cash_capex', '经营现金流出': 'operating_outflows', '经营现金流入': 'operating_inflows',
    '期间月数': 'months', '期初权益': 'equity_open', '期初应收账款': 'ar_open', '期初存货': 'inventory_open',
    '期初应付账款': 'ap_open', '赊销收入': 'credit_revenue', '赊购额': 'credit_purchases', '期间天数': 'period_days',
    '有息债务': 'interest_bearing_debt', '现金流量表现金': 'cfs_cash', '其他现金口径差异': 'other_cash_scope_difference',
}
SCALES = {'元': 1, '万元': 10000, '亿元': 100000000, 'RMB': 1, 'CNY': 1, '月': 1, '天': 1}


def record(label, value, unit, meta, source, status='READ', formula=None, cached=None):
    original = value
    try:
        if status not in ('READ', 'RECOMPUTED', 'REVIEWED'): raise ValueError('待复核或未支持')
        if unit not in SCALES: raise ValueError('未知单位')
        number = dec(value) * SCALES[unit]
    except ValueError:
        number = None
        if status == 'READ': status = 'MISSING_OR_INVALID'
    return dict(record_id=fingerprint({'source': source, 'label': label})[:20], original_account=label,
                original_value=original, original_unit=unit, currency=meta.get('币种'),
                period=meta.get('期间'), entity=meta.get('主体'), scope=meta.get('范围'), version=meta.get('版本'),
                cutoff=meta.get('截止日'), standard_account=ALIASES.get(label), value=number,
                mapping_rule='exact alias + explicit scale', source=source, status=status,
                formula=formula, cached_value=cached, independent_value=number if status == 'RECOMPUTED' else None)


def limited_formula(wb, sheet, cell, visiting=None):
    visiting = set() if visiting is None else visiting
    key = (sheet, cell)
    if key in visiting: raise ValueError('循环引用')
    visiting = visiting | {key}
    v = wb[sheet][cell].value
    if not isinstance(v, str) or not v.startswith('='):
        return dec(v)
    expr = v[1:].replace('$', '')
    if any(x in expr for x in ('[', ']', '!', '"', "'")):
        raise ValueError('v0.1不支持跨表/外链/字符串公式')
    # Only same-sheet SUM rectangular ranges and arithmetic cell refs.
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    def sum_range(m):
        a,b,c,d = range_boundaries(m.group(1))
        if (c-a+1)*(d-b+1) > 10000: raise ValueError('公式范围过大')
        return str(sum((limited_formula(wb, sheet, f'{get_column_letter(col)}{row}', visiting)
                        for row in range(b,d+1) for col in range(a,c+1)), dec(0)))
    expr = re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)', sum_range, expr, flags=re.I)
    def reference(m):
        return '(' + str(limited_formula(wb, sheet, m.group(0).upper(), visiting)) + ')'
    expr = re.sub(r'\b[A-Za-z]+\d+\b', reference, expr)
    try:
        tree = ast.parse(expr, mode='eval')
    except SyntaxError as exc:
        raise ValueError('v0.1不支持该Excel公式语法') from exc
    def compute(n):
        if isinstance(n, ast.Expression): return compute(n.body)
        if isinstance(n, ast.Constant) and type(n.value) in (int, float):
            return dec(ast.get_source_segment(expr, n))
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.UAdd, ast.USub)):
            return compute(n.operand) * (-1 if isinstance(n.op, ast.USub) else 1)
        if isinstance(n, ast.BinOp) and isinstance(n.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            a,b = compute(n.left), compute(n.right)
            if isinstance(n.op, ast.Add): return a+b
            if isinstance(n.op, ast.Sub): return a-b
            if isinstance(n.op, ast.Mult): return a*b
            if b == 0: raise ValueError('公式除零')
            return a/b
        raise ValueError('v0.1不支持该Excel公式')
    return compute(tree)


def extract_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    digest = file_hash(path)
    from layouts import matrix_records
    records, handled = matrix_records(wb, cached, path, digest)
    raw, warnings = [], []
    for ws in wb:
        meta = {}
        for row in ws.iter_rows():
            if len(row) >= 2 and row[0].value in ('主体','期间','范围','币种','版本','截止日','单位'):
                meta[row[0].value] = str(row[1].value)
            for c in row:
                if c.value is not None:
                    raw.append(dict(sheet=ws.title, cell=c.coordinate, value=c.value, number_format=c.number_format,
                                    hidden_row=bool(ws.row_dimensions[c.row].hidden),
                                    hidden_column=bool(ws.column_dimensions[c.column_letter].hidden),
                                    formula=c.value if c.data_type == 'f' else None,
                                    cached_value=cached[ws.title][c.coordinate].value if c.data_type == 'f' else None))
        for row in ws.iter_rows(min_row=1):
            if ws.title in handled or len(row) < 2 or row[0].value not in ALIASES: continue
            label, c = row[0].value, row[1]
            unit = str(row[2].value) if len(row)>2 and row[2].value else meta.get('单位')
            value, status, formula = c.value, 'READ', None
            if c.data_type == 'f':
                formula = c.value
                try: value, status = limited_formula(wb, ws.title, c.coordinate), 'RECOMPUTED'
                except (ValueError, SyntaxError, ArithmeticError, KeyError) as exc:
                    value, status = None, 'UNSUPPORTED_FORMULA'
                    warnings.append(f'{ws.title}!{c.coordinate}: {exc}')
            source = dict(file=str(Path(path).resolve()), sha256=digest, sheet=ws.title, cell=c.coordinate,
                          sheet_state=ws.sheet_state, merged_ranges=[str(r) for r in ws.merged_cells.ranges])
            rec = record(label, value, unit, meta, source, status, formula, cached[ws.title][c.coordinate].value if formula else None)
            rec['original_value'] = c.value
            records.append(rec)
    wb.close(); cached.close()
    return dict(status='EXTRACTED' if records else 'NEEDS_MAPPING', records=records, raw=raw,
                warnings=warnings, sha256=digest,
                limitation='自动映射仅支持已列科目的两列/三列报表；其余原始单元格保留，需补充映射，不要求用户整理JSON')


def parse_text_rows(lines, path, digest, page, ocr=False, boxes=None):
    meta = {}
    for line in lines:
        m = re.match(r'^(主体|期间|范围|币种|版本|截止日|单位)\s*[:：]\s*(.+)$', line.strip())
        if m: meta[m[1]] = m[2].strip()
    records = []
    for i,line in enumerate(lines):
        for label in ALIASES:
            match = re.match(r'^' + re.escape(label) + r'\s*(?:[|：:]\s*|\s+)([-−(（\d,，.)）]+)\s*(?:\|\s*)?(万元|亿元|元|月|天)?\s*$', line.strip())
            if not match: continue
            source = dict(file=str(Path(path).resolve()), sha256=digest, physical_page=page,
                          printed_page=None, line=i+1, original_line=line, bbox=(boxes or {}).get(i+1),
                          bbox_status='OCR_TOKEN_BOUNDS' if ocr else 'LINE_BASELINE_CANDIDATE_VERIFY_ON_PAGE')
            records.append(record(label, match[1].replace('（','(').replace('）',')'), match[2] or meta.get('单位'), meta,
                                  source, 'NEEDS_REVIEW' if ocr else 'READ'))
    return records


def extract_pdf(path):
    import pdfplumber
    digest = file_hash(path)
    from layouts import continued_text
    pages, records, inherited = [], [], {}
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ''
            words = page.extract_words()
            lines = text.splitlines()
            boxes = {}
            # Line bboxes retained from actual text layout, not invented page coordinates.
            for i, line in enumerate(lines, 1):
                # Coordinates of all words at corresponding distinct baseline, when ordering agrees.
                tops = sorted(set(round(w['top'], 1) for w in words))
                if i <= len(tops):
                    ws = [w for w in words if round(w['top'],1) == tops[i-1]]
                    boxes[i] = [min(w['x0'] for w in ws), min(w['top'] for w in ws), max(w['x1'] for w in ws), max(w['bottom'] for w in ws)]
            pages.append({'physical_page':n, 'text':text, 'words':words, 'needs_ocr':not bool(text.strip())})
            found,inherited = continued_text(lines,path,digest,n,inherited,boxes)
            records += found
    return dict(status='EXTRACTED' if records else 'NEEDS_OCR_OR_MAPPING', records=records, raw=pages, sha256=digest,
                limitation='复杂/跨页表格保留原页和词坐标；不保证全自动映射。扫描PDF需先在批准临时目录渲染，再用本地OCR。')


def extract_image(path):
    # System binary, no remote OCR. Output to stdout, not default cache directories.
    proc = subprocess.run(['/opt/homebrew/bin/tesseract', str(Path(path).resolve()), 'stdout', '-l', 'chi_sim+eng', '--psm','6','tsv'],
                          capture_output=True, text=True, timeout=45, check=True)
    rows = list(csv.DictReader(io.StringIO(proc.stdout), delimiter='\t'))
    grouped = {}
    for r in rows:
        if r.get('text','').strip():
            grouped.setdefault((r['block_num'],r['par_num'],r['line_num']),[]).append(r)
    lines, boxes = [], {}
    for ws in grouped.values():
        line = ' '.join(w['text'] for w in ws)
        # OCR separates Chinese glyphs; joining only between CJK characters preserves numeric tokens.
        line = re.sub(r'(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])','',line)
        lines.append(line)
        boxes[len(lines)] = [min(int(w['left']) for w in ws), min(int(w['top']) for w in ws),
                             max(int(w['left'])+int(w['width']) for w in ws), max(int(w['top'])+int(w['height']) for w in ws)]
    digest = file_hash(path)
    return dict(status='NEEDS_REVIEW', records=parse_text_rows(lines,path,digest,1,True,boxes),
                raw={'text':'\n'.join(lines),'tokens':rows}, sha256=digest,
                limitation='全部OCR候选值须核对原图；置信度不是准确性证明')


def extract(path):
    p = Path(path)
    before = file_hash(p)
    if p.suffix.lower() == '.xlsx': output = extract_xlsx(p)
    elif p.suffix.lower() == '.pdf': output = extract_pdf(p)
    elif p.suffix.lower() in ('.png','.jpg','.jpeg'): output = extract_image(p)
    else:
        output = dict(status='NOT_SUPPORTED', records=[], reason='v0.1不执行宏或自动转换旧XLS；保留原件，请在获准环境提供XLSX副本')
    if file_hash(p) != before: raise RuntimeError('原始材料被修改')
    output['original_unchanged'] = True
    return output


def review_record(r, value, reviewer, evidence):
    if not reviewer or not evidence: raise ValueError('缺少人工确认依据')
    new = dict(r)
    new['value'] = dec(value) * SCALES[r['original_unit']]
    new['status'] = 'REVIEWED'
    new['review'] = dict(reviewer=reviewer, evidence=evidence, previous_value=r['original_value'], confirmed_raw=value)
    return new


def metric_inputs(records, entity, scope, period, currency='CNY'):
    data, sources, blocked = {}, {}, {}
    for r in records:
        key = r.get('standard_account')
        if not key: continue
        if (r.get('entity'),r.get('scope'),r.get('period'),r.get('currency')) != (entity,scope,period,currency):
            continue
        if key in sources:
            data.pop(key,None); blocked[key] = '同口径重复/冲突，须选择版本'
            continue
        sources[key] = r['record_id']
        if r['status'] not in ('READ','RECOMPUTED','REVIEWED') or r['value'] is None:
            blocked[key] = r['status']
        else: data[key] = r['value']
    return data, sources, blocked
