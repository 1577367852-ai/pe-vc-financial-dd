"""Auditable snapshot tables; fresh result bundles require source+code+scenario fingerprints."""
import csv
import json
import math
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from common import ROOT, allowed_write, code_hash, dec, file_hash, fingerprint, safe_text, serial, write_json, write_text

LABELS = dict(zip(['gross_profit','gross_margin','net_margin','revenue_growth','current_ratio','debt_to_assets','rd_expense_share','ocf_margin','free_cash_flow','operating_nwc','dso','dio','dpo','roe','dscr','gross_operating_burn','net_operating_burn','pre_financing_burn','simple_runway'],
                 ['毛利（元）','毛利率','净利率','收入增长率','流动比率','总负债率','研发费用占经营费用比例','经营现金流/收入','简化自由现金流（元）','经营性营运资本（元）','应收周转天数','存货周转天数','应付周转天数','ROE（未年化）','偿债覆盖倍数','月经营总消耗（元）','月经营净消耗（元）','月融资前消耗含Capex（元）','简化现金跑道（月）']))
HEADERS = {'date':'日期','opening':'期初现金','inflows':'现金流入','outflows':'现金流出','closing':'期末现金','minimum':'最低现金要求','gap_zero':'零现金底线缺口','gap_floor':'含底线缺口',
           'operating_net_burn':'经营净消耗','financing_inflow':'增资流入','event_ids':'现金事件编号','name':'指标/检查','status':'状态','value':'独立计算结果','formula':'公式','inputs':'采用输入','reason':'说明',
           'source_record_ids':'来源记录编号','tolerance':'容差','id':'编号','amount':'金额','direction':'收付方向','category':'类别','source':'依据','origin':'来源类型',
           'record_id':'记录编号','original_account':'原始科目','original_value':'原始数值','original_unit':'原始单位','entity':'主体','scope':'范围','period':'期间','version':'版本','currency':'币种','cutoff':'截止日',
           'cached_value':'Excel缓存值','independent_value':'独立复算值','source_location':'原件位置','file_name':'原文件','adjustment_id':'调整编号','bucket':'调整类别'}
HEADERS.update({'finding_id':'问题编号','kind':'证据类型','fact':'具体发现','impact':'判断影响','request':'需补充材料',
                'base':'调整基数','adjustment_total':'已采纳调整合计','included_ids':'采纳编号','excluded_ids':'未采纳编号',
                'base_account':'基数科目','base_source':'基数口径依据','reviewer':'确认人','review_evidence':'确认依据',
                'source_sha256':'原件指纹','standard_account':'标准科目','confirmed_raw':'确认的原单位数值',
                'evidence_file':'证据文件','evidence_location':'证据位置','economic_item_id':'经济事项编号','approved':'是否采纳'})


def display_metric(name, value):
    if value is None:return '不可计算'
    if name in ('gross_margin','net_margin','revenue_growth','debt_to_assets','rd_expense_share','ocf_margin','roe'):
        return f'{value*100:.1f}%'
    return f'{value:,.2f}'


def dependency_signature(paths, assumptions):
    return fingerprint({'files':{str(Path(p).resolve()):file_hash(p) for p in paths}, 'assumptions':assumptions, 'code':code_hash()})


def assert_fresh(folder, paths, assumptions):
    f = Path(folder)
    m = json.loads((f/'manifest.json').read_text())
    expected = dependency_signature(paths, assumptions)
    if m['status'] != 'READY' or m['signature'] != expected:
        raise ValueError('STALE: 输入、合同、假设或代码变化；旧结果禁止继续引用，必须重算')
    for name, digest in m['outputs'].items():
        if not (f/name).exists() or file_hash(f/name) != digest:
            raise ValueError('STALE: 输出被修改或缺失；摘要和表格需一并重算')
    return True


def flatten_rows(rows):
    keys = list(dict.fromkeys(k for r in rows for k in r))
    return keys, [[serial(r.get(k)) for k in keys] for r in rows]


def readable(v):
    if isinstance(v, (dict,list,tuple)): return json.dumps(serial(v), ensure_ascii=False)
    if isinstance(v, Decimal): return float(v)
    return v


def write_csv(path, rows):
    p = allowed_write(path); p.parent.mkdir(parents=True,exist_ok=True)
    keys, vals = flatten_rows(rows)
    with p.open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(keys)
        for row in vals:w.writerow([safe_text(readable(v)) for v in row])


def workbook(path, tables, signature):
    # User-approved openpyxl authoring. Never execute source formulas/macros.
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from ingest import limited_formula
    wb=Workbook();wb.remove(wb.active)
    for name, rows in tables.items():
        ws=wb.create_sheet(name[:31]);ws.sheet_view.showGridLines=False
        keys=list(dict.fromkeys(k for r in rows for k in r)) or ['说明']
        ws.append([name+'（独立计算快照，修改输入后须重新运行）'])
        ws.append(['结果指纹', signature]);ws.append([HEADERS.get(k,k) for k in keys])
        for r in rows:
            values=[]
            for k in keys:
                value=readable(r.get(k))
                if value is not None and k in ('amount','base','value','adjustment_total','opening','inflows','outflows','closing','minimum','gap_zero','gap_floor','operating_net_burn','financing_inflow'):
                    try:value=float(dec(value))
                    except ValueError:pass
                if isinstance(value,str) and k in ('date','cutoff','baseline','accrual_start','accrual_end','截止日','基准日'):
                    try:value=datetime.fromisoformat(value)
                    except ValueError:pass
                values.append(safe_text(value))
            ws.append(values)
        ws.freeze_panes='B4';ws.auto_filter.ref=f'A3:{get_column_letter(len(keys))}{max(3,ws.max_row)}'
        for c in ws[3]:c.fill=PatternFill('solid',fgColor='29475F');c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
        ws['A1'].font=Font(name='Arial',bold=True,size=13)
        for row in ws.iter_rows(min_row=4):
            for c in row:
                c.font=Font(name='Arial',size=10)
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,(int,float)):c.number_format='#,##0.00;(#,##0.00);0.00'
                elif isinstance(c.value,datetime):c.number_format='yyyy-mm-dd'
                if name=='指标与校验' and c.column==keys.index('value')+1 and ws.cell(c.row,1).value in [LABELS[x] for x in ('gross_margin','net_margin','revenue_growth','debt_to_assets','rd_expense_share','ocf_margin','roe')]:
                    c.number_format='0.0%'
        for col,key in enumerate(keys,1):
            ws.column_dimensions[get_column_letter(col)].width=(18 if key in ('closing','opening','inflows','outflows','minimum','value','amount','base','adjustment_total')
                else 48 if key in ('fact','impact','request','source','inputs','formula','reason','review_evidence','base_source') else 28)
        for row in range(4,ws.max_row+1):
            line_count=1
            for c in ws[row]:
                text='' if c.value is None else str(c.value)
                visual_len=sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in text)
                width=ws.column_dimensions[c.column_letter].width or 18
                line_count=max(line_count,math.ceil(visual_len/max(6,width-2)))
            ws.row_dimensions[row].height=max(28,14*line_count+8)
        # Useful independent arithmetic, not cached source formula claims.
        if name=='现金预测':
            cols={k:get_column_letter(i+1) for i,k in enumerate(keys)}
            cidx=len(keys)+1;ws.cell(3,cidx,'Excel算术复核')
            for row in range(4,ws.max_row+1):
                ws.cell(row,cidx,f'={cols["opening"]}{row}+{cols["inflows"]}{row}-{cols["outflows"]}{row}-{cols["closing"]}{row}')
                if limited_formula(wb,ws.title,f'{get_column_letter(cidx)}{row}') != 0:raise ValueError('现金输出算术不一致')
            ws.column_dimensions[get_column_letter(cidx)].width=20
            ws.cell(3,cidx).fill=PatternFill('solid',fgColor='29475F');ws.cell(3,cidx).font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    p=allowed_write(path);p.parent.mkdir(parents=True,exist_ok=True);wb.save(p);wb.close()
    reopened=load_workbook(p,data_only=False,keep_links=False)
    if not reopened.sheetnames:raise ValueError('导出工作簿为空')
    if '现金预测' in reopened:
        ws=reopened['现金预测']
        for row in range(4,ws.max_row+1):
            if limited_formula(reopened,ws.title,f'{get_column_letter(ws.max_column)}{row}') != 0:raise ValueError('保存后公式复核失败')
    reopened.close()


def summary(report, mode):
    cash=report.get('cash',{})
    lines=[f'# {report["name"]}：'+('学习模式' if mode=='learning' else '工作模式'),'',
           '仅使用合成样本。计算及流程通过不代表财务资料真实，也不代表项目可投资。', '',
           f'财务截止日：{report["cutoff"]}；分析基准日：{report["baseline"]}。',
           f'期初现金状态：{report["opening"]["status"]}；口径：{report["opening"].get("label",report["opening"].get("reason",""))}。',
           f'该口径现金金额：{report["opening"].get("value")}；金额对应日期：{report["opening"].get("cash_as_of")}。',
           '所有金额为人民币元；本文件为一次运行快照。请通过view入口读取；离线打开文件不会自动检查过期状态。', '']
    if cash.get('status')=='PASS':
        lines += [f'- 情景：{cash["scenario"]["name"]}。',
                  f'- 第一次低于最低现金要求：{cash["first_warning"] or "预测期内未出现"}。',
                  f'- 第一次现金耗尽：{cash["depletion"] or "预测期内未出现"}。',
                  f'- 首次负余额：{cash["first_negative"] or "预测期内未出现"}。',
                  f'- 预测期最大零现金底线缺口：{cash["max_zero_gap"]}；含最低现金要求的缺口：{cash["max_floor_gap"]}。',
                  '- 该缺口按最早可注资的简化口径测算；分期投资仍需核对每个付款时点。',
                  '- 应核实付款节点、融资条件和收入回款。净资产或融资到账不能证明经营已停止烧钱。','']
    else: lines += ['现金预测未完成：'+cash.get('reason','缺输入'),'']
    lines += ['## 财务指标与状态','', '|指标|结果|状态|','|---|---:|---|']
    for k,v in report['metrics'].items():lines.append(f'|{LABELS.get(k,k)}|{display_metric(k,v["value"])}|{v["status"]}|')
    lines += ['', '历史平均消耗和未来预算分开计算：本示例现金预测来自独立预算，不能把历史月均值直接当作未来支出。历史指标的缺项不会被补零。']
    lines += ['', '## 调整口径','']
    for bucket,r in report.get('bridges',{}).items():
        lines.append(f'- {bucket}：基数{r.get("base")}；调整后{r.get("value")}；状态{r["status"]}。{r.get("reason","")}')
    lines += ['', '## 需进一步核实的事项','']
    for finding in report.get('findings',[]):
        lines += [f'### {finding["finding_id"]}','',finding['fact'],
                  f'涉及金额：{finding["amount"] if finding["amount"] is not None else "尚不能量化"}。',
                  '影响：'+finding['impact'],'需要材料：'+finding['request'],
                  '依据：'+json.dumps(serial(finding['source']),ensure_ascii=False),'']
    if not report.get('findings'):lines.append('本次未生成专项发现，不代表已排除风险；请查看指标状态及证据缺项。')
    if mode=='learning':
        lines += ['', '## 如何复核','', '先看原始与标准化数据表的来源位置，再看计算表的公式与输入。',
                  '经营净Burn =（经营现金流出－经营现金流入）/期间月数。融资属于筹资，不抵减经营Burn。',
                  '现金耗尽为余额≤0；首次负余额为余额<0；最低现金预警为余额低于设定底线。',
                  '合同总额、已收现金和收入确认不是同一口径。OCR候选数字在人工核对前不进入确定计算。']
    return '\n'.join(lines)+'\n'


def export_bundle(folder, report, paths, assumptions, expected_signature=None):
    f=allowed_write(folder);f.mkdir(parents=True,exist_ok=True)
    sig=dependency_signature(paths,assumptions)
    if expected_signature is not None and sig!=expected_signature:
        raise ValueError('运行期间输入已变化，禁止发布混合版本结果')
    write_json(f/'manifest.json',{'status':'RECOMPUTING','signature':sig})
    source_rows=[]
    for r in report['records']:
        src=r['source']
        source_rows.append({**{k:r.get(k) for k in ('record_id','original_account','original_value','original_unit','value','status','entity','scope','period','version','currency','cutoff','formula','cached_value','independent_value')},
                            'file_name':Path(src['file']).name,'source_location':(f"{src['sheet']}!{src.get('cell','')}" if src.get('sheet') else f"物理页{src.get('physical_page','图像')} 行{src.get('line','')} 坐标{src.get('bbox','')}"),
                            'source_sha256':src.get('sha256'), 'standard_account':r.get('standard_account'),
                            'reviewer':r.get('review',{}).get('reviewer'), 'review_evidence':r.get('review',{}).get('evidence'),
                            'confirmed_raw':r.get('review',{}).get('confirmed_raw')})
    calc_rows=[]
    from financial import CHECKS
    for k,v in list(report['metrics'].items())+list(report['checks'].items()):
        shown=dict(v)
        if k in CHECKS:shown['inputs']={key:v.get('inputs',{}).get(key) for key in CHECKS[k]}
        calc_rows.append(dict(name=LABELS.get(k,k),**shown))
    adjustment_rows=[]
    for item in report.get('adjustments',[]):
        adjustment_rows.append({**{k:item.get(k) for k in ('adjustment_id','economic_item_id','bucket','amount','status','reason','reviewer','approved')},
            'evidence_file':Path(item.get('evidence_file','')).name,'evidence_location':item.get('evidence_location'),
            'source_sha256':(item.get('source') or {}).get('sha256') if isinstance(item.get('source'),dict) else None})
    bridge_rows=[]
    for k,v in report.get('bridges',{}).items():
        bridge_rows.append(dict(bucket=k,base=v.get('base'),adjustment_total=(v['value']-v['base'] if v['status']=='PASS' else None),
            value=v.get('value'),status=v['status'],included_ids=[x['adjustment_id'] for x in v.get('inputs',{}).get('items',[])],
            excluded_ids=[x['adjustment_id'] for x in v.get('excluded',[])],base_account=v.get('base_account'),base_source=v.get('base_source'),reason=v.get('reason')))
    tables={
        '摘要':[{'项目':report['name'],'截止日':report['cutoff'],'基准日':report['baseline'],'现金状态':report['opening']['status'],'边界':'合成样本；非审计；请通过view检查是否过期'}],
        '现金预测':report.get('cash',{}).get('rows',[]),
        '指标与校验':calc_rows,
        '合同与现金事件':report.get('cash',{}).get('schedule',{}).get('events',[]),
        '原始与标准化数据':source_rows,
        '调整与问题':adjustment_rows or [{'说明':'未提供调整事项，不表示不存在应调整事项。'}],
        '调整桥':bridge_rows,
        '尽调问题':report.get('findings',[]),
    }
    write_json(f/'result.json',dict(report,signature=sig))
    write_text(f/'工作摘要.md',summary(report,'work'))
    write_text(f/'学习说明.md',summary(report,'learning'))
    for name,rows in tables.items():write_csv(f/(name+'.csv'),rows)
    workbook(f/'财务分析.xlsx',tables,sig)
    outputs=[p for p in f.iterdir() if p.is_file() and p.name!='manifest.json']
    # Recheck all input dependencies before claiming completion.
    if dependency_signature(paths,assumptions)!=sig:raise ValueError('运行期间输入已变化；结果仍为RECOMPUTING')
    write_json(f/'manifest.json',dict(status='READY',signature=sig,code=code_hash(),
               generated_utc=datetime.now(timezone.utc).isoformat(),
               inputs={str(Path(p).resolve()):file_hash(p) for p in paths}, assumptions=assumptions,
               outputs={p.name:file_hash(p) for p in outputs}))
    assert_fresh(f,paths,assumptions)
    return f
